import json, os, sys, subprocess
sys.path.insert(0, os.path.dirname(__file__))
import content_ko as K
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = "/home/claude/work"; OUT = f"{ROOT}/out"
FB = {"img_b2_depart": "img_b2_prep", "img_f15e": "img_spruance", "img_f16_fly": "img_b2_prep", "img_hornet": "img_ford_deck", "img_tehran_smoke": "img_borujerdi4"}
MAN = json.load(open(f"{ROOT}/media/manifest.json")); DUR = json.load(open(f"{ROOT}/audio/durations.json"))
F = "Arial"; HF = Font(name=F, bold=True, color="FFFFFF"); HFILL = PatternFill("solid", fgColor="1D3A5F")
thin = Side(style="thin", color="BBBBBB"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
FIX = PatternFill("solid", fgColor="FDECEA"); NEW = PatternFill("solid", fgColor="E8F3EA"); KEEP = PatternFill("solid", fgColor="F2F2F2")

def mk(k): return k if k in MAN else FB.get(k, k)

def probe(f):
    if not os.path.exists(f): return None
    j = json.loads(subprocess.check_output(["ffprobe", "-v", "error", "-show_entries", "format=duration,size:stream=codec_type,width,height,codec_name",
                                            "-of", "json", f]))
    v = next((s for s in j.get("streams", []) if s["codec_type"] == "video"), {})
    return dict(w=v.get("width"), h=v.get("height"), dur=round(float(j["format"]["duration"]), 1), mb=round(int(j["format"]["size"]) / 1e6, 1))

def sheet(wb, title, headers, rows, widths, fills=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]: c.font = HF; c.fill = HFILL; c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = B
    for i, r in enumerate(rows):
        ws.append(r)
        for c in ws[ws.max_row]:
            c.font = Font(name=F, size=10); c.alignment = Alignment(vertical="top", wrap_text=True); c.border = B
            if fills and fills[i]: c.fill = fills[i]
    for i, w in enumerate(widths): ws.column_dimensions[get_column_letter(i + 1)].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    return ws

OLD = {  # probed from the delivered zip
 "longform": ("epic_fury_ep01_longform_ko.mp4.html", "영상 아님 (HTML 44 KB)", "-"),
 "x": ("epic_fury_ep01_x_ko.mp4.mp4", "1920×1080", 61.9), "facebook": ("epic_fury_ep01_facebook_ko.mp4.mp4", "1080×1080", 55.7),
 "shorts": ("epic_fury_ep01_shorts_ko.mp4.mp4", "1080×1920", 91.9), "reels": ("epic_fury_ep01_instareels_ko.mp4.mp4", "1080×1920", 60.9),
 "tiktok": ("epic_fury_ep01_tiktok_ko.mp4.mp4", "1080×1920", 80.5), "threads": ("epic_fury_ep01_threads_ko.mp4.mp4", "1080×1920", 77.1),
 "feeds": ("(없음)", "-", "-"),
}
NAME = {"longform": "Longform (YouTube)", "x": "X", "facebook": "Facebook", "shorts": "YouTube Shorts", "reels": "Instagram Reels",
        "tiktok": "TikTok", "threads": "Threads", "feeds": "Instagram Feeds"}
REQ = {"longform": "3840×2160 (8분 이상)", "x": "3840×2160", "facebook": "3840×2160", "shorts": "2160×3840 (50초~2분 30초)",
       "reels": "2160×3840", "tiktok": "2160×3840", "threads": "2160×2700", "feeds": "2160×2700"}

def main():
    wb = Workbook(); wb.remove(wb.active)
    new = {p: probe(f"{OUT}/epic_fury_ep01_{p}_ko.mp4") for p in K.PLATFORMS}

    # ---- 1. change list ----
    R = []; Fl = []
    def add(scope, cat, before, basis, after, result, sev="높음"):
        R.append([len(R) + 1, scope, cat, sev, before, basis, after, result]); Fl.append({"수정": FIX, "신규": NEW, "유지": KEEP}.get(result.split()[0], None))
    allv = "전체 (영상 6개 + 롱폼 HTML)"
    add(allv, "내용 (줄거리)", "세종대왕함·박 함장·유조선 한빛호·서울 국가위기관리센터·광화문 편의점·동맹국 화상회의 등 옛 줄거리로 제작. ep01_final.html에는 이 인물·장소가 하나도 없음(검색 0건).",
        "epic_fury_ep01_final.html 본문(장면 1~8: 아자드 교수·펜타곤·다리우스 대령·웹 함장/새라·김준호 대위·레일라·나리만·포드 항모 전단)",
        "모든 영상의 대본·화면 문장을 HTML 본문 문장으로 다시 구성(요약 발췌, 창작 문장 추가 없음). 나레이션도 화면 문장 그대로 낭독.", "수정 (전면 재제작)")
    add(allv, "제목", "'제 1화 - 성벽이 무너진 날' / 표지 '작전명, 장대한 분노' 등", "HTML 제목 '제1화 · 성벽이 무너지는 날'",
        "에피소드 띠를 '제1화 · 성벽이 무너지는 날'로 통일. '제1화' 붙여 쓰기(접두사 '제').", "수정")
    add(allv, "사실관계 (시각)", "워싱턴 D.C. 02:00 대통령 승인", "HTML 타임라인: 2월 27일 15:38(미 동부) 실행 명령",
        "n1 '2월 27일 오후 3시 38분, 미 동부 시각 … 실행을 명령했다'로 교체", "수정")
    add(allv, "사실관계 (시각)", "테헤란 14:17 지휘부 직격, 최고지도자 사망 몇 시간 뒤 공식화", "HTML: 10:15 관저 단지 타격, 이란 정부 공식 확인 3월 1일",
        "n3·h_1015 '오전 10시 15분 … 공식 확인은 3월 1일'", "수정")
    add(allv, "사실관계 (보복)", "16:00 샤하브-3 계열 탄도미사일로 바레인·카타르·이라크 미군 기지 공격", "HTML: 12:35 '진실의 약속 4' 개시, 12:56 바레인 미 5함대 시설 피격, 걸프·아르빌·요르단 확대",
        "n5 문장으로 교체(샤하브-3 언급 삭제)", "수정")
    add(allv, "사실관계 (위치)", "'같은 시각 호르무즈 해협, 항공모함 포드함의 갑판 위로 드론들이…'", "HTML: 포드 항모 전단은 동지중해 이스라엘 앞바다(지도 밖)",
        "포드함 장면은 모두 '동지중해'로 표기. 호르무즈 해협은 해도 화면·역사 메모로만 등장", "수정")
    add(allv, "사실관계 (속보 띠)", "'이란 최고지도자 아야톨라 알리 하메네이 사망' 단정, 문장이 잘림('(Ali…')", "HTML 속보: 이스라엘 발표, 이란 공식 확인은 3월 1일",
        "속보 띠를 HTML 속보 5건으로 교체, 끊김 없이 흐르는 스크롤 띠로 변경", "수정")
    add(allv, "날짜 표기", "롱폼 HTML 6~13쪽 머리 날짜 '2026. 2. 26'", "개전일 2026년 2월 28일", "모든 화면 머리 날짜 '2026. 2. 28'", "수정", "중간")
    add(allv, "숫자·시각 표기 (korean-style-guide)", "'오전 아홉 시 사십오 분', '오후 두 시 십칠 분' 등 한글 수사", "HTML은 아라비아 숫자('오전 9시 45분')",
        "HTML 표기로 통일. 작전명은 작은따옴표, 대사는 큰따옴표로 통일", "수정", "낮음")
    add(allv, "쪽 번호", "영상 'x/21', HTML '1/13' 불일치", "-", "플랫폼별 실제 쪽수로 'n/N' 자동 표기", "수정", "낮음")
    add(allv, "실존 인물·픽션 고지", "픽션 고지 없음", "HTML 머리말: 실제 작전에서 착상한 픽션, 인물·대화는 창작",
        "모든 표지에 '실제 전쟁에서 착상한 소설입니다…' 표기, 롱폼 첫 쪽에 작가의 말 낭독, 엔딩에 'AI 내레이션(Typecast 인화)' 표기. 가상 사건(에빈 구 오폭)·가상 감청 문서에는 '소설 속 가상 사건/FICTION' 꼬리표", "수정")
    add(allv, "배경 영상·출처", "지휘센터·유조선·구축함·서울 전광판 스톡 영상, 화면에 출처 없음", "규칙: 무료 출처 4K, 이미지 바로 아래 출처, 전쟁 이후 자료 우선",
        "개전 이후 실사(Wikimedia Commons: 스프루언스 토마호크 발사 2/28, 포드함 비행갑판 3/1, B-2, 테헤란 피격 주택가 3/4·3/6, 이란 미사일·드론 발사 지도)를 우선 쓰고, 장소를 특정하지 않는 분위기 영상만 Pexels 4K로 보충. 모든 박스 아래 출처(작가·라이선스·사이트) 표기", "수정")
    add(allv, "레이아웃 (프로젝트 규칙)", "16:9 영상도 텍스트를 이미지 위에 겹쳐 배치", "왼쪽 텍스트 박스 + 오른쪽 16:9 이미지 박스, 이미지는 텍스트 박스 세로 중앙, 텍스트는 왼쪽 경계에서 5mm",
        "가로형(롱폼·X·Facebook)은 좌 텍스트/우 16:9 박스(같은 높이·세로 중앙), 세로형·피드형은 16:9 박스 아래 텍스트 박스. 글자는 박스 왼쪽 경계에서 38px(4K 기준 5mm)", "수정", "중간")
    add(allv, "동적 효과", "정지 화면 위주, 일부 페이지 하단이 검게 비어 있음(Reels 18~19쪽)", "'다이내믹 버전'",
        "배경 블러·감광 + 박스 영상 켄 번스 이동, 글자 0.5초 선행 페이드·슬라이드 인, 속보 띠 스크롤, 표지 서서히 확대", "수정", "중간")
    add(allv, "나레이션", "옛 대본 낭독", "KR = Typecast '인화'", "Typecast 인화(voice_id tc_6296a815b958a8ed9610b189, ssfm-v30)로 전 문장 새로 생성, loudnorm −16 LUFS", "수정")
    add(allv, "파일 이름", "'.mp4.mp4' 이중 확장자, 롱폼은 '.mp4.html'", "epic fury epxx [platform] [ko]", "epic_fury_ep01_{platform}_ko.mp4 (+ .srt 자막)", "수정", "낮음")
    for p in K.PLATFORMS:
        o = OLD[p]; n = new[p] or {}
        if p == "longform":
            add(NAME[p], "형식", "롱폼 영상 파일 없음 — zip 안의 파일은 HTML(44 KB)", REQ[p], f"새 MP4 {n.get('w')}×{n.get('h')}, {n.get('dur')}초", "신규")
        elif p == "feeds":
            add(NAME[p], "누락", "zip에 Instagram Feeds 영상 없음", REQ[p], f"새 MP4 {n.get('w')}×{n.get('h')}, {n.get('dur')}초", "신규")
        else:
            add(NAME[p], "해상도", f"{o[1]} (4K 미만)" + (" · 정사각형 비율" if p == "facebook" else " · 9:16 (규격은 4:5)" if p == "threads" else ""),
                REQ[p], f"{n.get('w')}×{n.get('h')}로 재렌더링", "수정")
    add(allv, "미디어 대체", "-", "Commons 원본 우선",
        "Wikimedia 서버가 이 작업 환경의 내려받기를 반복 차단(HTTP 429)해 B-2 이륙·F-15E·F-16·F/A-18E·테헤란 북부 피격 사진 5장은 받지 못함. 같은 범주의 개전 이후 Commons 사진(B-2 출격 준비, 스프루언스, 포드함, 보루제르디 주택가)으로 대체. 원하면 PC에서 받아 GitHub imgs/novel/ep01에 올린 뒤 같은 렌더러로 해당 쪽만 다시 굽기 가능", "유지 (대체)", "낮음")
    add("EN·JP 판", "범위 밖", "zip에 EN·JP 영상 없음(검토 대상 아님)", "원 지시: ko·en·jp", "이번 수정 범위에서 제외. ElevenLabs 잔여 291자로 EN·JP 나레이션 생성 불가 — 크레딧 충전 후 같은 렌더러로 제작 가능", "유지 (보류)", "참고")
    ws = sheet(wb, "수정 사항 목록", ["No", "대상", "구분", "중요도", "기존(문제)", "기준", "수정 내용", "결과"], R, [5, 22, 18, 8, 50, 42, 60, 14], Fl)

    # ---- 2. spec compare ----
    rows = []
    for p in K.PLATFORMS:
        o = OLD[p]; n = new[p] or {}
        rows.append([NAME[p], o[0], o[1], o[2], REQ[p], f"epic_fury_ep01_{p}_ko.mp4", f"{n.get('w')}×{n.get('h')}", n.get("dur"), n.get("mb"),
                     K.COVERS[p]["title"], len(K.PLATFORMS[p][2]) + 1])
    sheet(wb, "파일 스펙 비교", ["플랫폼", "기존 파일", "기존 해상도", "기존 길이(초)", "요구 규격", "수정 파일", "수정 해상도", "수정 길이(초)", "용량(MB)", "표지(썸네일) 제목", "쪽 수"],
          rows, [18, 36, 22, 12, 22, 34, 14, 12, 10, 28, 8])

    # ---- 3. scripts ----
    rows = []
    for p, (W, H, seq, lay) in K.PLATFORMS.items():
        rows.append([NAME[p], 0, "cover_" + p, "표지", K.COVERS[p]["narr"], DUR["cover_" + p]["dur"], mk(K.COVERS[p]["media"])])
        for i, pid in enumerate(seq, 1):
            d = K.P.get(pid) or K.C[pid]
            rows.append([NAME[p], i, pid, d["label"], d["text"], DUR[pid]["dur"], mk(d["media"])])
    sheet(wb, "대본(화면=낭독)", ["플랫폼", "순서", "ID", "위치·시각 라벨", "화면 문장 = 나레이션", "음성 길이(초)", "미디어 키"], rows, [18, 6, 12, 34, 80, 10, 18])

    # ---- 4. media ----
    used = {}
    for p, (W, H, seq, lay) in K.PLATFORMS.items():
        for pid in ["cover"] + seq:
            m = K.COVERS[p]["media"] if pid == "cover" else (K.P.get(pid) or K.C[pid])["media"]
            used.setdefault(mk(m), set()).add(NAME[p])
    rows = []
    for k, m in sorted(MAN.items()):
        if k not in used and not any(k == v for v in []): continue
        rows.append([k, m["type"], m.get("title"), m.get("site"), m.get("author"), m.get("license"), m.get("date") or "", f"{m.get('width')}×{m.get('height')}",
                     m.get("page"), ", ".join(sorted(used.get(k, [])))])
    sheet(wb, "미디어 출처", ["키", "종류", "제목/검색어", "사이트", "작가", "라이선스", "촬영일", "원본 해상도", "원본 페이지", "사용 플랫폼"], rows, [18, 8, 40, 16, 26, 22, 18, 12, 60, 40])

    # ---- 5. tools ----
    T = [
     ["korean-style-guide (스킬)", "영상 문장 ↔ HTML 표기 일관성(숫자·시각·따옴표·인명)", "한글 수사 시각 → 아라비아 숫자, 따옴표 규칙 통일"],
     ["korean-grammar-checker (스킬)", "기존 화면 문장·새 대본 맞춤법·띄어쓰기", "'제 1화' → '제1화' 외 새 대본 오류 없음"],
     ["vidIQ (커넥터)", "유튜브 제목 클릭 점수", "롱폼 '장대한 분노 1화: 성벽이 무너지는 날 | 2026 이란 전쟁 소설' 88점 / '오전 9시 45분, 중동의 하늘이 찢어졌다 | 장대한 분노 1화 (전쟁 소설)' 87점 · 쇼츠 '폭탄이 200미터 빗나간 아침 | 장대한 분노 1화' 79점 / '작전명, 장대한 분노 | 전쟁 소설 1화' 78점"],
     ["Typecast API", "KR 나레이션 '인화'", f"{len(DUR)}개 문장 단위 음성 생성"],
     ["Pexels API · Wikimedia Commons", "4K 무료 영상·개전 이후 사진", "출처 시트 참조"],
     ["ffmpeg · Playwright", "4K 렌더링·화면 레이어", "8개 플랫폼 MP4 + SRT"],
     ["Firecrawl (미사용)", "사실 확인", "HTML을 기준 원고로 삼고 HTML 출처(Wikipedia 등)를 그대로 따름. 웹 작업 시 무관 커넥터를 쓰지 않는다는 저장된 규칙에 따라 생략"],
     ["humanize-scan · human-touch (미사용)", "본문 윤문", "영상 문장은 HTML 본문을 그대로 옮긴 것이라 윤문 대상 아님(PDF 5항: 사실 표기·본문 보존)"],
     ["adobe-for-creativity · Canva · postiz (미사용)", "크기 변환·예약 발행", "규격별 렌더를 직접 만들어 크기 변환 불필요. 발행은 이번 요청 범위 밖"],
    ]
    sheet(wb, "적용 스킬·도구", ["스킬·커넥터·플러그인", "쓰임", "결과"], T, [36, 40, 90])

    # ---- 0. summary (first) ----
    ws = wb.create_sheet("요약", 0)
    lines = [["장대한 분노 EP01 KR 플랫폼 영상 검토·수정 결과 (2026-10-27)"], [],
             ["기준 원고", "epic_fury_ep01_final.html"], ["검토 대상", "mp4.zip — 영상 6개(X·Facebook·Shorts·Reels·TikTok·Threads) + 롱폼 HTML 1개"],
             ["판정", "6개 영상 모두 기준 HTML과 줄거리·제목·시각이 다르고 4K 규격 미달 → 전면 재제작. 롱폼·Instagram Feeds는 영상이 없어 신규 제작."],
             ["수정 항목 수", f"=COUNTA('수정 사항 목록'!A2:A200)"],
             ["결과 영상", "8개 (4K, H.264/AAC, 30fps) + 자막 SRT 8개"],
             ["나레이션", "Typecast '인화' (ssfm-v30)"],
             ["범위 밖", "EN·JP 판 (zip에 없음, ElevenLabs 크레딧 부족)"]]
    for r in lines: ws.append(r)
    ws["A1"].font = Font(name=F, bold=True, size=14)
    for row in ws.iter_rows(min_row=3):
        for c in row: c.font = Font(name=F, size=11, bold=(c.column == 1)); c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 18; ws.column_dimensions["B"].width = 110
    wb.save(f"{OUT}/epic_fury_ep01_ko_video_review.xlsx")

main()
